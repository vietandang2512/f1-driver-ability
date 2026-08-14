"""
Builds the final deliverable spreadsheet: full 2014-present driver rankings
from all four models (race ELO, qualifying ELO, race ridge, qualifying
ridge), one combined sheet plus one full-detail sheet per model.
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DATA = Path(__file__).resolve().parent.parent / "data" / "processed"
OUT = Path(__file__).resolve().parent.parent / "F1_Driver_Rankings_2014_present.xlsx"

FONT = "Arial"
HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=13)
NOTE_FONT = Font(name=FONT, italic=True, size=9, color="555555")
BODY_FONT = Font(name=FONT, size=10)


def load_master():
    with open(DATA / "master_rankings_2014_present.csv", newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def driver_label(d):
    return d.replace("-", " ").title()


def build_combined_sheet(wb, master_rows):
    ws = wb.active
    ws.title = "Combined"

    ws["A1"] = "F1 Driver Ability — Combined Rankings, 2014–Present"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:I1")

    ws["A2"] = ("Two models (ridge regression, chained ELO) x two metrics (race pace, qualifying pace). "
                "Ridge coefficients are in milliseconds, negative = faster than average. "
                "ELO ratings are on a 1500-centered scale. Sorted by average rank across all four models.")
    ws["A2"].font = NOTE_FONT
    ws.merge_cells("A2:I2")

    headers = [
        "Driver",
        "Race ELO rank", "Race ELO rating",
        "Qualifying ELO rank", "Qualifying ELO rating",
        "Race ridge rank", "Race ridge (ms)",
        "Qualifying ridge rank", "Qualifying ridge (ms)",
        "Avg rank (available models)",
    ]
    header_row = 4
    for i, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=i, value=h)
    style_header(ws, header_row, len(headers))

    # Pre-sort by a simple python-computed average rank (for row ORDER only --
    # the displayed "avg rank" column itself is a live formula so it recalculates
    # if anyone edits ranks later).
    def avg_rank_key(row):
        ranks = [row[k] for k in ("race_elo_rank", "quali_elo_rank", "race_ridge_rank", "quali_ridge_rank") if row[k]]
        return sum(int(r) for r in ranks) / len(ranks) if ranks else 999

    master_rows_sorted = sorted(master_rows, key=avg_rank_key)

    r = header_row + 1
    for row in master_rows_sorted:
        ws.cell(row=r, column=1, value=driver_label(row["driver"]))
        ws.cell(row=r, column=2, value=int(row["race_elo_rank"]) if row["race_elo_rank"] else None)
        ws.cell(row=r, column=3, value=float(row["race_elo_rating"]) if row["race_elo_rating"] else None)
        ws.cell(row=r, column=4, value=int(row["quali_elo_rank"]) if row["quali_elo_rank"] else None)
        ws.cell(row=r, column=5, value=float(row["quali_elo_rating"]) if row["quali_elo_rating"] else None)
        ws.cell(row=r, column=6, value=int(row["race_ridge_rank"]) if row["race_ridge_rank"] else None)
        ws.cell(row=r, column=7, value=float(row["race_ridge_ms"]) if row["race_ridge_ms"] else None)
        ws.cell(row=r, column=8, value=int(row["quali_ridge_rank"]) if row["quali_ridge_rank"] else None)
        ws.cell(row=r, column=9, value=float(row["quali_ridge_ms"]) if row["quali_ridge_ms"] else None)
        # live formula: average of whichever rank cells in this row are populated
        ws.cell(row=r, column=10, value=f"=AVERAGE(B{r},D{r},F{r},H{r})")
        for c in range(1, 11):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            if c in (3, 5):
                cell.number_format = "0.0"
            if c in (7, 9):
                cell.number_format = "#,##0.0"
            if c == 10:
                cell.number_format = "0.00"
        r += 1

    ws.freeze_panes = "B5"
    autosize(ws, [20, 12, 13, 16, 16, 13, 15, 16, 17, 20])
    return r - 1  # last data row


def build_model_sheet(wb, title, master_rows, value_key, rank_key, n_key, is_ridge):
    ws = wb.create_sheet(title=title)
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")

    metric_label = "coefficient (ms, negative = faster)" if is_ridge else "rating (1500-centered)"
    headers = ["Rank", "Driver", metric_label, "Sample size"]
    header_row = 3
    for i, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=i, value=h)
    style_header(ws, header_row, len(headers))

    rows = [row for row in master_rows if row[rank_key]]
    rows_sorted = sorted(rows, key=lambda row: int(row[rank_key]))

    r = header_row + 1
    for row in rows_sorted:
        ws.cell(row=r, column=1, value=int(row[rank_key]))
        ws.cell(row=r, column=2, value=driver_label(row["driver"]))
        ws.cell(row=r, column=3, value=float(row[value_key]))
        ws.cell(row=r, column=4, value=int(row[n_key]))
        for c in range(1, 5):
            ws.cell(row=r, column=c).font = BODY_FONT
        ws.cell(row=r, column=3).number_format = "#,##0.0" if is_ridge else "0.0"
        r += 1

    ws.freeze_panes = "A4"
    autosize(ws, [7, 20, 26, 13])


if __name__ == "__main__":
    master_rows = load_master()

    wb = Workbook()
    build_combined_sheet(wb, master_rows)

    build_model_sheet(wb, "Race ELO (full)", master_rows,
                       "race_elo_rating", "race_elo_rank", "race_elo_n", is_ridge=False)
    build_model_sheet(wb, "Qualifying ELO (full)", master_rows,
                       "quali_elo_rating", "quali_elo_rank", "quali_elo_n", is_ridge=False)
    build_model_sheet(wb, "Race ridge (full)", master_rows,
                       "race_ridge_ms", "race_ridge_rank", "race_ridge_n", is_ridge=True)
    build_model_sheet(wb, "Qualifying ridge (full)", master_rows,
                       "quali_ridge_ms", "quali_ridge_rank", "quali_ridge_n", is_ridge=True)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote workbook -> {OUT}")
